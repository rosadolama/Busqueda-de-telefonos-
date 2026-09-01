"""Une varios Excels (o CSV) de contactos de empresas en uno solo, fusionando
los contactos repetidos entre archivos.

Pensado para el caso real: tienes varios listados sueltos de la misma
empresa —uno con email, otro con teléfono, otro con redes sociales— y
algunos contactos aparecen en más de un archivo. Este script no hace nada
más que eso: juntar todo en un único Excel, una fila por empresa/contacto,
con todos los datos combinados.

Uso:
  python dedup_contactos.py contactos1.xlsx contactos2.xlsx contactos3.xlsx
  python dedup_contactos.py --carpeta ./mis_excels -o unificado.xlsx

Qué hace exactamente:
  1. Lee todas las filas de todos los archivos indicados (.xlsx/.xlsm/.csv).
     Cada archivo puede tener columnas distintas (uno solo trae email, otro
     solo teléfono...) y con nombres distintos ("Correo", "Mail", "Tel",
     "Móvil"...): se reconocen automáticamente las variantes habituales.
  2. Detecta qué filas son el mismo contacto repetido entre archivos —por
     email, dominio de la web, o nombre de empresa— y las fusiona en una
     sola fila combinando todos los datos disponibles.
  3. Escribe un Excel nuevo con una fila por contacto/empresa, más una
     columna "Fuentes" (de qué archivo(s) viene) y "Revisar" para los pocos
     casos en que la fusión automática no tiene datos suficientes para
     decidir sola.

No inventa ni corrige datos: si dos fuentes traen valores distintos para el
mismo campo (dos teléfonos distintos, por ejemplo), ambos se conservan en la
misma celda separados por "; " en vez de descartar uno a ciegas.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

CANONICAL_FIELDS = [
    "Empresa", "Nombre", "Apellidos", "Cargo", "Email", "Teléfono",
    "Web", "LinkedIn", "Instagram", "Facebook", "Twitter/X",
    "Redes sociales", "Dirección",
]

# Claves ya normalizadas (sin tildes, en minúsculas, sin palabras de enlace
# como "de/la/del": ver _normalize_header) para reconocer cabeceras
# equivalentes aunque cada archivo las llame distinto.
ALIASES: dict[str, str] = {
    "empresa": "Empresa",
    "compania": "Empresa",
    "company": "Empresa",
    "nombre empresa": "Empresa",
    "empresa nombre": "Empresa",
    "razon social": "Empresa",
    "nombre comercial": "Empresa",
    "negocio": "Empresa",
    "title": "Empresa",  # nombre de campo típico en exports de Google Maps/Apify
    "nombre": "Nombre",
    "contacto": "Nombre",
    "nombre contacto": "Nombre",
    "persona contacto": "Nombre",
    "persona": "Nombre",
    "name": "Nombre",
    "first name": "Nombre",
    "apellidos": "Apellidos",
    "apellido": "Apellidos",
    "surname": "Apellidos",
    "last name": "Apellidos",
    "cargo": "Cargo",
    "puesto": "Cargo",
    "posicion": "Cargo",
    "role": "Cargo",
    "position": "Cargo",
    "email": "Email",
    "e mail": "Email",
    "correo": "Email",
    "correo electronico": "Email",
    "mail": "Email",
    "telefono": "Teléfono",
    "tel": "Teléfono",
    "movil": "Teléfono",
    "celular": "Teléfono",
    "phone": "Teléfono",
    "whatsapp": "Teléfono",
    "numero telefono": "Teléfono",
    "numero contacto": "Teléfono",
    "web": "Web",
    "website": "Web",
    "pagina web": "Web",
    "sitio web": "Web",
    "url": "Web",
    "linkedin": "LinkedIn",
    "instagram": "Instagram",
    "ig": "Instagram",
    "facebook": "Facebook",
    "fb": "Facebook",
    "twitter": "Twitter/X",
    "red social": "Redes sociales",
    "redes sociales": "Redes sociales",
    "redes": "Redes sociales",
    "social media": "Redes sociales",
    "social": "Redes sociales",
    "direccion": "Dirección",
    "domicilio": "Dirección",
    "address": "Dirección",
    "ubicacion": "Dirección",
    "localidad": "Dirección",
    "poblacion": "Dirección",
    "ciudad": "Dirección",
    "city": "Dirección",
}
# Alias multi-palabra primero, para que "correo electronico" no se quede
# matcheando solo "correo" y deje colgando "electronico" sin usar.
_SORTED_ALIASES = sorted(ALIASES.items(), key=lambda kv: -len(kv[0].split(" ")))

_STOPWORDS = {"de", "del", "la", "las", "los", "el", "un", "una"}

# Para estos campos, una palabra que sobra en la cabecera puede cambiar el
# significado por completo en vez de ser un simple matiz: "Verificada WEB"
# no es la web (es un booleano sobre ella), "Nombre Oportunidad" no es el
# nombre del contacto (es el nombre de un trato de CRM), "Seguidores IG" o
# "Mandado Instagram" no son el perfil de Instagram (son métricas sobre él).
# Solo se acepta el match si TODO lo que sobra está en esta lista de
# calificadores conocidos; si no, se sigue buscando otro alias.
_RESTRICTED_QUALIFIERS: dict[str, set[str]] = {
    # set() = no se acepta ninguna palabra de sobra: "empresa" suelto solo
    # cuenta si es la cabecera entera (o ya viene consumido del todo por un
    # alias de dos palabras como "nombre empresa"). Sin esto, "empresa" --
    # insertado primero en ALIASES -- le roba el match a alias de una sola
    # palabra más específicos como "linkedin" en cabeceras tipo "LinkedIn
    # Empresa", porque los empates de longitud se prueban en orden de
    # inserción del diccionario.
    "Empresa": set(),
    "Web": {"url", "link", "sitio", "pagina", "empresa"},
    "LinkedIn": {"url", "link", "perfil", "profile", "pagina", "empresa"},
    "Instagram": {"url", "link", "perfil", "profile", "pagina", "empresa"},
    "Facebook": {"url", "link", "perfil", "profile", "pagina", "empresa"},
    "Twitter/X": {"url", "link", "perfil", "profile", "pagina", "empresa"},
    "Nombre": {"contacto", "decisor", "responsable", "encargado"},
}


_CAMEL_RE = re.compile(r"(?<=[a-z0-9])(?=[A-Z])")


def _normalize_header(h) -> str:
    if h is None:
        return ""
    text = str(h).strip()
    # exports tipo Apify/API aplanan JSON anidado en camelCase y con "/"
    # como separador de ruta ("linkedinUrl", "affiliatedPages/0/name"):
    # separar el camelCase en palabras antes de perder las mayúsculas, si
    # no "linkedinUrl" queda pegado como una sola palabra irreconocible.
    text = _CAMEL_RE.sub(" ", text).lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    tokens = [t for t in text.split() if t and t not in _STOPWORDS]
    return " ".join(tokens)


def _match_header(header_norm: str) -> str | None:
    if not header_norm:
        return None
    tokens = header_norm.split(" ")
    for alias, canonical in _SORTED_ALIASES:
        alias_tokens = alias.split(" ")
        n = len(alias_tokens)
        if n > len(tokens):
            continue
        for i in range(len(tokens) - n + 1):
            if tokens[i:i + n] != alias_tokens:
                continue
            qualifiers = _RESTRICTED_QUALIFIERS.get(canonical)
            if qualifiers is not None:
                resto = tokens[:i] + tokens[i + n:]
                if any(t not in qualifiers for t in resto):
                    continue  # la palabra que sobra cambia el sentido: no es este campo
            return canonical
    return None


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_header_row(rows: list[tuple]) -> int:
    """Busca en las primeras filas cuál es la cabecera real: a veces la fila
    1 es un título o va en blanco. Se queda con la fila (de las 5 primeras)
    que más columnas reconocidas tiene; si ninguna reconoce nada, usa la
    primera fila no vacía tal cual."""
    best_idx, best_score = 0, -1
    first_nonempty = None
    for i, row in enumerate(rows[:5]):
        texts = [_cell_to_str(v) for v in row]
        if first_nonempty is None and any(texts):
            first_nonempty = i
        score = sum(1 for t in texts if t and _match_header(_normalize_header(t)))
        if score > best_score:
            best_score, best_idx = score, i
    if best_score <= 0:
        return first_nonempty if first_nonempty is not None else 0
    return best_idx


def _resolve_field_map(header_row) -> tuple[dict[int, str], dict[int, str]]:
    field_map: dict[int, str] = {}
    extra_headers: dict[int, str] = {}
    for idx, h in enumerate(header_row):
        text = _cell_to_str(h)
        if not text:
            continue
        canonical = _match_header(_normalize_header(text))
        if canonical:
            field_map[idx] = canonical
        else:
            extra_headers[idx] = text

    # Ambigüedad real: una columna suelta llamada "Nombre" en un listado de
    # CONTACTOS DE EMPRESAS suele ser el nombre de la empresa, salvo que ya
    # haya Empresa/Apellidos/Cargo en la misma tabla -- ahí sí es claramente
    # el nombre de pila de la persona de contacto.
    mapped = set(field_map.values())
    if "Nombre" in mapped and not mapped & {"Empresa", "Apellidos", "Cargo"}:
        for idx, canonical in field_map.items():
            if canonical == "Nombre":
                field_map[idx] = "Empresa"

    return field_map, extra_headers


def _print_mapping_report(origen: str, header_row, field_map: dict, extra_headers: dict) -> None:
    partes = [f"{canonical}<-'{_cell_to_str(header_row[idx])}'" for idx, canonical in field_map.items()]
    linea = f"  {origen}: " + (", ".join(partes) if partes else "(ninguna columna reconocida)")
    if extra_headers:
        linea += f" | sin reconocer: {list(extra_headers.values())}"
    print(linea)


def _looks_like_placeholder_row(campos: dict) -> bool:
    """Algunos exports de CRM/scraper (Apify, GoHighLevel...) dejan una fila
    justo debajo de la cabecera con los NOMBRES de los campos como si fueran
    valores reales (p.ej. Web='website', Teléfono='phone', Empresa='title').
    Se detecta de forma genérica: si el valor de una celda, pasado por el
    mismo reconocedor de cabeceras, apunta a ESE MISMO campo, es eco del
    nombre de su propia columna, no un dato real. Con un solo eco podría ser
    casualidad; con dos o más ya es un patrón."""
    echoes = sum(
        1 for campo, valor in campos.items()
        if valor and _match_header(_normalize_header(valor)) == campo
    )
    return echoes >= 2


def _build_records(header_row, data_rows, field_map: dict, extra_headers: dict) -> list[tuple[dict, dict]]:
    records = []
    for raw_row in data_rows:
        campos = {f: "" for f in CANONICAL_FIELDS}
        extra: dict[str, str] = {}
        has_canonical_value = False
        for idx, value in enumerate(raw_row):
            text = _cell_to_str(value)
            if not text:
                continue
            if idx in field_map:
                campo = field_map[idx]
                campos[campo] = f"{campos[campo]}; {text}" if campos[campo] else text
                has_canonical_value = True
            elif idx in extra_headers:
                header = extra_headers[idx]
                extra[header] = f"{extra[header]}; {text}" if header in extra else text
        # Se exige al menos un campo CANÓNICO (no solo columnas sin
        # reconocer) para contar como contacto real: descarta filas de hojas
        # que en realidad son texto explicativo/leyenda, no datos.
        if has_canonical_value and not _looks_like_placeholder_row(campos):
            records.append((campos, extra))
    return records


def _read_table(rows: list[tuple], origen_label: str) -> list[tuple[dict, dict]]:
    if not rows:
        return []
    header_idx = _find_header_row(rows)
    header_row = rows[header_idx]
    field_map, extra_headers = _resolve_field_map(header_row)
    if not field_map and not extra_headers:
        return []
    _print_mapping_report(origen_label, header_row, field_map, extra_headers)
    return _build_records(header_row, rows[header_idx + 1:], field_map, extra_headers)


def read_xlsx_file(path: Path, hoja: str | None = None) -> list[tuple[dict, dict]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    records = []
    multi_sheet = len(wb.worksheets) > 1
    hoja_norm = hoja.strip().lower() if hoja else None
    nombres_hojas = [ws.title for ws in wb.worksheets]
    for ws in wb.worksheets:
        if hoja_norm is not None and ws.title.strip().lower() != hoja_norm:
            continue
        rows = list(ws.iter_rows(values_only=True))
        label = f"{path.name} [{ws.title}]" if multi_sheet else path.name
        records += _read_table(rows, label)
    wb.close()
    if hoja_norm is not None and hoja_norm not in (n.strip().lower() for n in nombres_hojas):
        raise ValueError(f"la hoja '{hoja}' no existe en {path.name} (hojas disponibles: {', '.join(nombres_hojas)})")
    return records


def read_csv_file(path: Path) -> list[tuple[dict, dict]]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(text.splitlines(), dialect))
    return _read_table(rows, path.name)


def read_file(path: Path, hoja: str | None = None) -> list[tuple[dict, dict]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx_file(path, hoja)
    if suffix == ".csv":
        return read_csv_file(path)
    raise ValueError(f"formato no soportado: {suffix}")


@dataclass
class Contacto:
    campos: dict = field(default_factory=dict)
    extra: dict = field(default_factory=dict)
    fuentes: set = field(default_factory=set)
    revisar: list = field(default_factory=list)
    n_fusionados: int = 1


def _first(value: str) -> str:
    return value.split(";")[0].strip() if value else ""


_SUFFIX_RE = re.compile(
    r"\b("
    r"s\.?\s?l\.?\s?u\.?"          # S.L.U.
    r"|s\.?\s?a\.?\s?u\.?"          # S.A.U.
    r"|s\.?\s?a\.?\s?t\.?"          # S.A.T.
    r"|s\.?\s?c\.?\s?a\.?"          # S.C.A.
    r"|s\.?\s?a\.?\s?s\.?"          # S.A.S.
    r"|s\.?\s?l\.?"                 # S.L.
    r"|s\.?\s?a\.?"                 # S.A.
    r"|soc\.?\s?coop\.?"
    r"|sdad\.?\s?coop\.?"
    r"|cooperativa"
    r"|coop\.?"
    r"|c\.?b\.?"
    r"|inc\.?"
    r"|llc\.?"
    r"|ltd\.?"
    r"|corp\.?"
    r")\b\.?",
    re.IGNORECASE,
)


def normalize_name(name: str) -> str:
    """minúsculas, sin tildes, sin formas societarias (S.L., S.A., Inc...),
    para que 'Acme S.L.' y 'ACME, S.L.U.' se reconozcan como la misma
    empresa al deduplicar."""
    if not name:
        return ""
    n = unicodedata.normalize("NFKD", name)
    n = "".join(c for c in n if not unicodedata.combining(c))
    n = n.lower()
    n = _SUFFIX_RE.sub(" ", n)
    n = _SUFFIX_RE.sub(" ", n)
    n = re.sub(r"[^a-z0-9 ]", " ", n)
    n = re.sub(r"\s+", " ", n).strip()
    return n


def normalize_email(email: str) -> str:
    return email.strip().lower() if email else ""


def base_domain(url: str) -> str:
    if not url:
        return ""
    netloc = urlparse(url if "//" in url else f"//{url}").netloc.lower()
    if netloc.startswith("www."):
        netloc = netloc[4:]
    return netloc.split(":")[0]


_SHARED_PLATFORM_DOMAINS = {
    # proveedores de email gratuitos: dos empresas distintas pueden usar
    # ambas @gmail.com, así que coincidir aquí no dice nada.
    "gmail.com", "hotmail.com", "outlook.com", "yahoo.com", "yahoo.es",
    "icloud.com", "live.com", "aol.com", "protonmail.com", "hotmail.es",
    # dominios de redes sociales/enlaces cortos: en listados sacados de
    # Google Maps es habitual que el campo "Web" lleve en realidad un
    # enlace de Instagram/Facebook/WhatsApp cuando la empresa no tiene web
    # propia -- tratarlo como "dominio de la empresa" fusionaría a ciegas
    # empresas distintas que solo comparten la plataforma.
    "facebook.com", "instagram.com", "linkedin.com", "twitter.com", "x.com",
    "wa.me", "whatsapp.com", "youtube.com", "tiktok.com", "linktr.ee",
    "goo.gl", "bit.ly", "maps.google.com", "g.page",
}


def domain_of(campos: dict) -> str:
    """Dominio identificativo de la empresa: prioriza la Web; si no hay,
    usa el dominio del email. En ambos casos descarta dominios compartidos
    (gmail.com, instagram.com...) donde coincidir no dice nada de verdad."""
    web_domain = base_domain(_first(campos.get("Web", "")))
    if web_domain and web_domain not in _SHARED_PLATFORM_DOMAINS:
        return web_domain
    email = _first(campos.get("Email", ""))
    if "@" in email:
        email_domain = email.split("@")[-1].strip().lower()
        if email_domain and email_domain not in _SHARED_PLATFORM_DOMAINS:
            return email_domain
    return ""


def normalize_phone(phone: str) -> str:
    if not phone:
        return ""
    digits = re.sub(r"\D", "", phone)
    if not digits:
        return ""
    return digits[-9:] if len(digits) > 9 else digits


def dedup_contactos(contactos: list[Contacto]) -> list[Contacto]:
    """Fusión transitiva (union-find) por email, dominio, nombre de empresa
    y teléfono -- misma idea que un dedup de directorio de empresas clásico:
    el email y el dominio son señales fuertes (fusionan solas); el nombre
    fusiona salvo conflicto de dominio; el teléfono es la señal más débil y
    solo fusiona si ya viene corroborado por dominio o nombre -- si no, se
    marca para revisar en vez de fusionar a ciegas (una centralita puede
    compartirse entre empresas distintas)."""
    n = len(contactos)
    parent = list(range(n))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def union(i: int, j: int) -> None:
        ri, rj = find(i), find(j)
        if ri != rj:
            parent[ri] = rj

    emails = [normalize_email(_first(c.campos.get("Email", ""))) for c in contactos]
    domains = [domain_of(c.campos) for c in contactos]
    names = [normalize_name(c.campos.get("Empresa", "")) for c in contactos]
    phones = [normalize_phone(_first(c.campos.get("Teléfono", ""))) for c in contactos]

    by_email: dict[str, list[int]] = {}
    for i, e in enumerate(emails):
        if e:
            by_email.setdefault(e, []).append(i)
    for idxs in by_email.values():
        for j in idxs[1:]:
            union(idxs[0], j)

    by_domain: dict[str, list[int]] = {}
    for i, d in enumerate(domains):
        if d:
            by_domain.setdefault(d, []).append(i)
    for idxs in by_domain.values():
        for j in idxs[1:]:
            union(idxs[0], j)

    by_name: dict[str, list[int]] = {}
    for i, nm in enumerate(names):
        if nm:
            by_name.setdefault(nm, []).append(i)
    for idxs in by_name.values():
        for a in range(len(idxs)):
            for b in range(a + 1, len(idxs)):
                i, j = idxs[a], idxs[b]
                if domains[i] and domains[j] and domains[i] != domains[j]:
                    continue  # mismo nombre pero dominios distintos -> empresas distintas
                union(i, j)

    by_phone: dict[str, list[int]] = {}
    for i, p in enumerate(phones):
        if p:
            by_phone.setdefault(p, []).append(i)
    for idxs in by_phone.values():
        for a in range(len(idxs)):
            for b in range(a + 1, len(idxs)):
                i, j = idxs[a], idxs[b]
                if find(i) == find(j):
                    continue  # ya fusionados por email/dominio/nombre
                same_domain = bool(domains[i]) and domains[i] == domains[j]
                same_name = bool(names[i]) and names[i] == names[j]
                if same_domain or same_name:
                    union(i, j)
                    continue
                emp_i = contactos[i].campos.get("Empresa") or "(sin empresa)"
                emp_j = contactos[j].campos.get("Empresa") or "(sin empresa)"
                # Nombres, no "a"/"b": esos ya son los contadores de los bucles
                # de fuera, reutilizarlos los corrompería a mitad de bucle.
                menor, mayor = sorted([emp_i, emp_j])
                motivo = f"Mismo teléfono entre '{menor}' y '{mayor}', sin más datos que lo confirmen"
                contactos[i].revisar.append(motivo)
                contactos[j].revisar.append(motivo)

    groups: dict[int, list[int]] = {}
    order: list[int] = []
    for i in range(n):
        root = find(i)
        if root not in groups:
            groups[root] = []
            order.append(root)
        groups[root].append(i)

    return [_merge_group([contactos[i] for i in groups[root]]) for root in order]


def _merge_group(group: list[Contacto]) -> Contacto:
    group.sort(key=lambda c: sum(1 for v in c.campos.values() if v), reverse=True)
    best = group[0]
    for other in group[1:]:
        for campo in CANONICAL_FIELDS:
            val_best = best.campos.get(campo, "")
            val_other = other.campos.get(campo, "")
            if not val_other:
                continue
            if not val_best:
                best.campos[campo] = val_other
            elif campo == "Empresa":
                # Empresa es la clave de fusión: las variantes ("Acme S.L."
                # vs "ACME, S.L.U.") son ruido de formato, no datos distintos
                # -- se queda con la de la fila más completa, sin concatenar.
                continue
            elif val_other not in val_best.split("; "):
                best.campos[campo] = f"{val_best}; {val_other}"
        for header, val in other.extra.items():
            if not val:
                continue
            if header not in best.extra:
                best.extra[header] = val
            elif val not in best.extra[header].split("; "):
                best.extra[header] = f"{best.extra[header]}; {val}"
        best.fuentes |= other.fuentes
        best.revisar.extend(other.revisar)
    best.n_fusionados = len(group)
    if best.n_fusionados > 1:
        # El dominio fusiona sin pedir corroboración (señal fuerte), pero eso
        # puede juntar de más una cadena con varias sedes -- cada una con su
        # propio Instagram/Facebook -- bajo la misma web corporativa. No se
        # deshace la fusión (la empresa/dominio SÍ es el mismo), pero se avisa
        # para que se revise si conviene separarlas por sede.
        for campo in ("Instagram", "Facebook", "LinkedIn", "Twitter/X"):
            valores = {v.strip() for v in best.campos.get(campo, "").split(";") if v.strip()}
            if len(valores) > 1:
                best.revisar.append(
                    f"{campo}: combina {len(valores)} cuentas distintas al fusionar -- comprueba si son sedes/perfiles distintos"
                )
    return best


def write_xlsx(contactos: list[Contacto], path: Path) -> None:
    extra_headers = sorted({h for c in contactos for h in c.extra})
    columns = CANONICAL_FIELDS + extra_headers + ["Fuentes", "Duplicados_fusionados", "Revisar"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Contactos"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)

    for c in contactos:
        row = [c.campos.get(col, "") for col in CANONICAL_FIELDS]
        row += [c.extra.get(col, "") for col in extra_headers]
        row.append("; ".join(sorted(c.fuentes)))
        row.append(c.n_fusionados)
        row.append(" | ".join(dict.fromkeys(c.revisar)))
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.font = Font(name="Arial")

    for col_cells in ws.columns:
        length = max((len(str(cell.value)) for cell in col_cells if cell.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 50)

    wb.save(path)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="dedup_contactos.py",
        description="Une varios Excels/CSV de contactos de empresas en uno solo, fusionando los repetidos entre archivos.",
        epilog=(
            "Ejemplos:\n"
            "  python dedup_contactos.py contactos1.xlsx contactos2.xlsx contactos3.xlsx\n"
            "  python dedup_contactos.py --carpeta ./mis_excels -o unificado.xlsx\n"
            "  python dedup_contactos.py \"crm.xlsx:Estetica España\" \"otro.xlsx:Empresas RSS\"\n"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "archivos", nargs="*",
        help="Rutas a los Excels (.xlsx/.xlsm) o CSV a unir. Para usar solo una hoja "
        "concreta de un Excel: 'archivo.xlsx:Nombre de la hoja' (si no se indica, se leen todas)",
    )
    parser.add_argument("--carpeta", help="Carpeta de la que coger todos los .xlsx/.xlsm/.csv (alternativa a listar archivos; siempre lee todas las hojas)")
    parser.add_argument("-o", "--salida", default="contactos_unificados.xlsx", help="Ruta del Excel de salida (por defecto: contactos_unificados.xlsx)")
    return parser


def _parse_archivo_arg(raw: str) -> tuple[Path, str | None]:
    """'archivo.xlsx' o 'archivo.xlsx:Nombre de hoja'. Solo se trata lo que
    va tras los ':' como filtro de hoja si lo de antes es un archivo que
    existe de verdad (para no romper una ruta que ya trajera ':' por otro
    motivo)."""
    if ":" in raw:
        ruta, _, hoja = raw.rpartition(":")
        if ruta and Path(ruta).exists():
            return Path(ruta), hoja
    return Path(raw), None


def _collect_input_paths(args) -> list[tuple[Path, str | None]]:
    salida_resuelta = Path(args.salida).resolve()
    entradas: list[tuple[Path, str | None]] = [_parse_archivo_arg(a) for a in args.archivos]
    if args.carpeta:
        carpeta = Path(args.carpeta)
        entradas += [(p, None) for p in sorted(carpeta.glob("*")) if p.suffix.lower() in (".xlsx", ".xlsm", ".csv")]
    seen = set()
    unicas: list[tuple[Path, str | None]] = []
    for path, hoja in entradas:
        clave = (path.resolve(), (hoja or "").strip().lower())
        # Si la salida cae dentro de --carpeta (p.ej. al repetir la ejecución),
        # no releerla como si fuera un archivo de contactos más.
        if path.resolve() == salida_resuelta or clave in seen:
            continue
        seen.add(clave)
        unicas.append((path, hoja))
    return unicas


def main():
    args = build_arg_parser().parse_args()
    entradas = _collect_input_paths(args)
    if not entradas:
        print("No se ha indicado ningún archivo. Pasa rutas a Excels/CSV o usa --carpeta.", file=sys.stderr)
        sys.exit(1)

    print("=== LEYENDO ARCHIVOS (columnas detectadas) ===")
    contactos: list[Contacto] = []
    por_archivo: dict[str, int] = {}
    for path, hoja in entradas:
        if not path.exists():
            print(f"Aviso: no existe el archivo '{path}', se omite.", file=sys.stderr)
            continue
        try:
            registros = read_file(path, hoja)
        except Exception as exc:  # noqa: BLE001 - un archivo con problemas no debe tumbar el resto
            print(f"Aviso: no se pudo leer '{path}' ({exc}), se omite.", file=sys.stderr)
            continue
        etiqueta = f"{path.name}:{hoja}" if hoja else path.name
        por_archivo[etiqueta] = len(registros)
        for campos, extra in registros:
            contactos.append(Contacto(campos=campos, extra=extra, fuentes={etiqueta}))

    if not contactos:
        print("No se ha leído ninguna fila válida de los archivos indicados.", file=sys.stderr)
        sys.exit(1)

    total_crudo = len(contactos)
    deduped = dedup_contactos(contactos)
    write_xlsx(deduped, Path(args.salida))

    revisar = [c for c in deduped if c.revisar]
    print("\n=== RESUMEN ===")
    print("Filas leídas por archivo:")
    for nombre, n in por_archivo.items():
        print(f"  {nombre}: {n}")
    print(f"Total filas leídas: {total_crudo}")
    print(f"Filas tras fusionar duplicados: {len(deduped)}  (fusionadas de más: {total_crudo - len(deduped)})")
    if revisar:
        print(f"\nFilas marcadas para revisar a mano: {len(revisar)}")
        for c in revisar:
            nombre = c.campos.get("Empresa") or "(sin empresa)"
            print(f"  {nombre}: {' | '.join(dict.fromkeys(c.revisar))}")
    print(f"\nExcel unificado escrito en: {args.salida}")


if __name__ == "__main__":
    main()
